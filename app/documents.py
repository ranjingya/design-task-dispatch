"""从需求文档提取带来源的文字和图片。"""

import io
from contextlib import closing
import logging
import posixpath
import warnings
from xml.etree import ElementTree
from zipfile import ZipFile

import openpyxl
import pypdfium2 as pdfium
import xlrd
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


def read_document(data: bytes, suffix: str, source: str, material) -> None:
    """功能说明：解析需求文档并汇入本次材料。参数 data：文件内容；suffix：文件类型；source：来源名称；material：共享限额与收集器。返回值：无。"""
    logger.debug("读取需求文档：type=%s bytes=%d", suffix, len(data))
    if suffix == ".pdf":
        read_pdf(data, source, material)
    elif suffix == ".xls":
        read_xls(data, source, material)
    else:
        read_xlsx(data, source, material)


def read_xlsx(data: bytes, source: str, material) -> None:
    """功能说明：提取工作表单元格、批注和锚定图片。参数 data：工作簿字节；source：来源名称；material：材料收集器。返回值：无。"""
    config = material.config
    with ZipFile(io.BytesIO(data)) as archive:
        if sum(item.file_size for item in archive.infolist()) > config["max_unpacked_bytes"]:
            raise ValueError("Excel 解压后大小超过上限")
        relationships = {item.attrib["Id"]: item.attrib["Target"] for item in
                         ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))}
        active_paths = set()
        for sheet in ElementTree.fromstring(archive.read("xl/workbook.xml")).findall("{*}sheets/{*}sheet"):
            if config["skip_hidden_sheets"] and sheet.get("state", "visible") != "visible":
                continue
            relation = next(value for key, value in sheet.attrib.items() if key.endswith("}id"))
            target = relationships[relation]
            active_paths.add(target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target))
        # 加载前限制实际单元格数量，并检查可见页的公式缓存。
        cell_count = 0
        for name in archive.namelist():
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                with archive.open(name) as stream:
                    for _, cell in ElementTree.iterparse(stream, events=("end",)):
                        if cell.tag.endswith("}c"):
                            cell_count += 1
                            if cell_count > config["max_cells"]:
                                raise ValueError("Excel 单元格数量超过上限")
                            formula = cell.find("{*}f")
                            value = cell.find("{*}v")
                            if name in active_paths and formula is not None and (value is None or value.text is None):
                                material.issue(source, "存在没有缓存计算结果的公式，请保存计算结果后重新提交")
                            cell.clear()
        # 单元格图片和嵌入对象无法由 openpyxl 完整还原，必须提示材料缺口。
        if any(name.startswith(("xl/cellimages", "xl/richData/", "xl/embeddings/")) for name in archive.namelist()):
            material.issue(source, "包含无法完整读取的单元格图片或嵌入对象")
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
    try:
        if caught:
            material.issue(source, "部分 Excel 对象无法完整解析")
        visible = [sheet for sheet in workbook if not config["skip_hidden_sheets"] or sheet.sheet_state == "visible"]
        if len(visible) > config["max_sheets"]:
            material.issue(source, "工作表数量超过上限")
        hidden = len(workbook.worksheets) - len(visible)
        if hidden:
            material.add_text(source, f"按配置跳过 {hidden} 张隐藏工作表")
        cells = 0
        for sheet in visible[:config["max_sheets"]]:
            label = f"{source} / {sheet.title}"
            cells += sheet.max_row * sheet.max_column
            if cells > config["max_cells"]:
                material.issue(label, "单元格数量超过上限")
                break
            material.add_text(label, "合并范围：" + str(sheet.merged_cells))
            sheet_has_text = False
            for row in sheet.iter_rows():
                values = []
                for cell in row:
                    if cell.value is not None:
                        sheet_has_text = True
                        values.append(f"{cell.coordinate}={cell.value}")
                    if cell.comment:
                        values.append(f"{cell.coordinate}批注={cell.comment.text}")
                if values and not material.add_text(label, " | ".join(values)):
                    break
            for picture in sheet._images:
                anchor = getattr(picture.anchor, "_from", None)
                position = f"{get_column_letter(anchor.col + 1)}{anchor.row + 1}" if anchor else "浮动图片"
                try:
                    # 列标题和附近单元格随目录发送；必读判断限定标题及明确看图指引。
                    nearby, headings = [], []
                    if anchor:
                        row, col = anchor.row + 1, anchor.col + 1
                        for nearby_row in sheet.iter_rows(min_row=max(1, row - 1), max_row=min(sheet.max_row, row + 1),
                                                    min_col=max(1, col - 1), max_col=min(sheet.max_column, col + 1)):
                            nearby.extend(f"{cell.coordinate}={cell.value}" for cell in nearby_row if cell.value is not None)
                        headings = [str(sheet.cell(r, col).value or "") for r in range(1, min(row, 4))]
                    context = " | ".join(headings + nearby)
                    visual_options = material.config["visual_review"]
                    required = (not sheet_has_text or any(word in sheet.title + " ".join(headings) for word in visual_options["instruction_keywords"])
                                or any(word in context for word in visual_options["image_reference_keywords"]))
                    material.add_image(picture._data(), f"{label} / {position}", group=label,
                                       context=context, requires_detail=required)
                except Exception:
                    material.issue(label, "内嵌图片无法读取")
            if sheet._charts:
                material.issue(label, "包含未渲染的图表")
    finally:
        workbook.close()


def read_xls(data: bytes, source: str, material) -> None:
    """功能说明：读取旧版 Excel 文本。参数 data：工作簿字节；source：来源名称；material：材料收集器。返回值：无。"""
    workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
    try:
        cells = 0
        sheets = [sheet for sheet in workbook.sheets() if not material.config["skip_hidden_sheets"] or not sheet.visibility]
        if len(sheets) > material.config["max_sheets"]:
            material.issue(source, "工作表数量超过上限")
        for sheet in sheets[:material.config["max_sheets"]]:
            cells += sheet.nrows * sheet.ncols
            if cells > material.config["max_cells"]:
                material.issue(source, "单元格数量超过上限")
                break
            for index in range(sheet.nrows):
                row = " | ".join(f"{get_column_letter(col + 1)}{index + 1}={value}" for col, value in enumerate(sheet.row_values(index)) if value != "")
                if row and not material.add_text(f"{source} / {sheet.name}", row):
                    break
        # XLS 解析器不提供图形覆盖信息，要求转为 XLSX 以确认完整材料。
        material.issue(source, "旧版 XLS 已读取文本；请转为 XLSX，以便核对图片与完整内容")
    finally:
        workbook.release_resources()


def read_pdf(data: bytes, source: str, material) -> None:
    """功能说明：逐页提取 PDF 文字并渲染页面供视觉识别。参数 data：PDF 字节；source：来源名称；material：材料收集器。返回值：无。"""
    with pdfium.PdfDocument(data) as document:
        if len(document) > material.config["max_pdf_pages"]:
            material.issue(source, "PDF 页数超过上限")
        for index in range(min(len(document), material.config["max_pdf_pages"])):
            label = f"{source} / 第 {index + 1} 页"
            with closing(document[index]) as page:
                with closing(page.get_textpage()) as textpage:
                    text = textpage.get_text_bounded()
                    material.add_text(label, text)
                if not material.image_available(label):
                    continue
                scale = min(2, material.config["visual_review"]["detail_tile_size"] / min(page.get_size()))
                if page.get_width() * page.get_height() * scale * scale > 40_000_000:
                    material.issue(label, "PDF 页面像素超过上限，请拆分长页面")
                    continue
                with closing(page.render(scale=scale)) as bitmap:
                    picture = bitmap.to_pil()
                    output = io.BytesIO()
                    picture.save(output, format="PNG")
                    visual_options = material.config["visual_review"]
                    required = not text.strip() or any(word in text for word in visual_options["instruction_keywords"] + visual_options["image_reference_keywords"])
                    material.add_image(output.getvalue(), label, group=source, context=text[:2000], requires_detail=required)
