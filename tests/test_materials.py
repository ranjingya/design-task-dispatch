"""验证真实文档解析、来源定位、限额与在线导出边界。"""

import io
import json
from dataclasses import replace
from unittest.mock import Mock

import openpyxl
import pytest
from PIL import Image
from openpyxl.drawing.image import Image as ExcelImage

from app.config import ROOT, Settings
from app.documents import read_document
from app.feishu import FeishuClient
from app.materials import MaterialBundle, collect_materials, sheet_token


@pytest.fixture
def settings(tmp_path):
    return Settings("app", "secret", "base", "table", "token", tmp_path / "fields.json")


@pytest.fixture
def config():
    return json.loads((ROOT / "config/materials.json").read_text())


def workbook_bytes(formula=False):
    """生成含需求、隐藏页和参考图片的小工作簿。"""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "修改要求"
    sheet["A1"] = "请制作商品主图"
    sheet["B2"] = "=1+1" if formula else "尺寸 750×750"
    sheet.merge_cells("A1:B1")
    image = io.BytesIO()
    Image.new("RGB", (100, 100), "red").save(image, format="PNG")
    sheet.add_image(ExcelImage(image), "C3")
    hidden = workbook.create_sheet("历史模板")
    hidden["A1"] = "不可用的旧需求"
    hidden.sheet_state = "hidden"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def pdf_bytes():
    """生成带文字页与扫描页的两页 PDF。"""
    import pypdfium2 as pdfium

    # 用 PDF 标准内容流构造文本页，第二页使用图像模拟扫描材料。
    content = b"BT /F1 20 Tf 30 150 Td (Create product image) Tj ET"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
               b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
               b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream"]
    data = b"%PDF-1.4\n"
    offsets = []
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += str(index).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    position = len(data)
    data += b"xref\n0 6\n0000000000 65535 f \n" + b"".join(f"{n:010d} 00000 n \n".encode() for n in offsets)
    data += b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n" + str(position).encode() + b"\n%%EOF"
    scan = io.BytesIO()
    Image.new("RGB", (200, 100), "blue").save(scan, format="PDF")
    with pdfium.PdfDocument(data) as document, pdfium.PdfDocument(scan.getvalue()) as scanned:
        document.import_pages(scanned)
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()


def test_excel_text_images_and_hidden_scope(settings, config):
    bundle = MaterialBundle(settings, config)
    read_document(workbook_bytes(), ".xlsx", "需求.xlsx", bundle)
    assert not bundle.issues
    assert any("A1=请制作商品主图" in row["text"] for row in bundle.texts)
    assert not any("不可用的旧需求" in row["text"] for row in bundle.texts)
    assert bundle.image_sources[0]["source"] == "需求.xlsx / 修改要求 / C3"
    assert bundle.image_sources[0]["requires_detail"] is True
    assert "750" in bundle.image_sources[0]["context"]
    assert len(bundle.images) == 1


def test_excel_missing_formula_cache_requires_review(settings, config):
    bundle = MaterialBundle(settings, config)
    read_document(workbook_bytes(True), ".xlsx", "需求.xlsx", bundle)
    assert any("公式" in issue for issue in bundle.issues)


def test_pdf_text_and_scanned_page(settings, config):
    bundle = MaterialBundle(settings, config)
    read_document(pdf_bytes(), ".pdf", "需求.pdf", bundle)
    assert not bundle.issues
    assert len(bundle.images) == 2
    assert any("Create product image" in row["text"] for row in bundle.texts)
    assert bundle.image_sources[1]["source"].endswith("第 2 页")


@pytest.mark.parametrize("limit", [{"max_pdf_pages": 1}, {"max_text_chars": 5}])
def test_pdf_partial_read_has_explicit_issue(settings, config, limit):
    bundle = MaterialBundle(settings, config | limit)
    read_document(pdf_bytes(), ".pdf", "需求.pdf", bundle)
    assert bundle.issues


def test_excel_cell_and_unpacked_limits(settings, config):
    for limit in ({"max_cells": 1}, {"max_unpacked_bytes": 1}):
        bundle = MaterialBundle(settings, config | limit)
        try:
            read_document(workbook_bytes(), ".xlsx", "需求.xlsx", bundle)
        except ValueError as exc:
            assert "上限" in str(exc)
        else:
            assert bundle.issues


def test_zip_skipped_and_bad_pdf_does_not_stop_other_materials(settings):
    feishu = Mock(spec=FeishuClient, settings=settings)
    feishu.download_attachment.side_effect = [b"broken pdf", workbook_bytes()]
    attachments = [{"file_token": "zip", "name": "图片.zip", "type": "application/zip"},
                   {"file_token": "pdf", "name": "需求.pdf"},
                   {"file_token": "excel", "name": "说明.xlsx"}]
    bundle = collect_materials({"fldTest": attachments}, None, "recTest", feishu)
    assert feishu.download_attachment.call_count == 2
    assert not bundle.issues
    assert len(bundle.skipped_attachments) == 1 and "需求.pdf" in bundle.skipped_attachments[0]['source']
    assert len(bundle.images) == 1


def test_unknown_attachments_do_not_consume_file_budget(settings, config, tmp_path, caplog):
    """多种未知格式不会阻止后续有效附件读取。"""
    path = tmp_path / 'materials.json'
    path.write_text(json.dumps(config | {'max_files': 1}))
    client = Mock(spec=FeishuClient, settings=replace(settings, materials_path=path))
    client.download_attachment.return_value = workbook_bytes()
    attachments = [{'file_token': 'psd', 'name': '模板.PSD'},
                   {'file_token': 'unknown', 'name': '素材.unknown'},
                   {'file_token': 'no_extension', 'name': '无后缀素材'},
                   {'file_token': 'excel', 'name': '需求.xlsx'}]
    bundle = collect_materials({'fldA': attachments}, None, 'recA', client)
    assert client.download_attachment.call_count == 1
    assert client.download_attachment.call_args.args[0] == 'excel'
    assert len(bundle.skipped_attachments) == 3 and not bundle.issues
    assert any('请制作商品主图' in row['text'] for row in bundle.texts)
    assert '跳过附件' in caplog.text and '模板.PSD' in caplog.text


@pytest.mark.parametrize('failure', [RuntimeError('下载权限不足'), OSError('读取失败'), b'broken image'])
def test_unreadable_attachment_keeps_following_valid_file(settings, failure):
    """下载失败或文件损坏仅排除当前附件。"""
    client = Mock(spec=FeishuClient, settings=settings)
    client.download_attachment.side_effect = [failure, workbook_bytes()]
    bundle = collect_materials({'fldA': [{'file_token': 'broken', 'name': '素材.png'},
                                        {'file_token': 'valid', 'name': '需求.xlsx'}]}, None, 'recA', client)
    assert len(bundle.skipped_attachments) == 1 and not bundle.issues
    assert len(bundle.images) == 1
    assert all('素材.png' not in item['source'] for item in bundle.image_sources)
    assert any('请制作商品主图' in row['text'] for row in bundle.texts)


@pytest.mark.parametrize('failure_mode', ['exception', 'incomplete'])
def test_failed_attachment_discards_partial_text_and_images(settings, monkeypatch, failure_mode):
    """失败附件的部分文字和图片不会混入后续匹配输入。"""
    client = Mock(spec=FeishuClient, settings=settings)
    client.download_attachment.side_effect = [b'partial', b'valid']

    def parse(data, suffix, source, bundle):
        """模拟附件解析中途失败或发现未读内容。"""
        if data == b'valid':
            bundle.add_text(source, '有效需求内容')
            return
        bundle.add_text(source, '失败附件的片段')
        image = io.BytesIO()
        Image.new('RGB', (20, 20), 'white').save(image, format='PNG')
        bundle.add_image(image.getvalue(), source)
        if failure_mode == 'exception':
            raise RuntimeError('文档后半部分无法解析')
        bundle.issue(source, '文档后半部分无法解析')

    monkeypatch.setattr('app.documents.read_document', parse)
    bundle = collect_materials({'fldA': [{'file_token': 'partial', 'name': '不完整.pdf'},
                                        {'file_token': 'valid', 'name': '需求.pdf'}]}, None, 'recA', client)
    assert [row['text'] for row in bundle.texts] == ['有效需求内容']
    assert not bundle.images and not bundle.image_sources and not bundle.image_parts
    assert bundle.image_size == 0
    assert bundle.text_size == sum(len(row['source']) + len(row['text']) for row in bundle.texts)
    assert len(bundle.skipped_attachments) == 1 and not bundle.issues
    assert bundle.file_size == len(b'partialvalid')


@pytest.mark.parametrize('limit', ['files', 'images'])
def test_attachment_fields_share_limits(settings, config, tmp_path, limit):
    """两列附件共用文件及图片限额，不因增加字段而翻倍。"""
    if limit == 'files':
        path = tmp_path / 'materials.json'
        path.write_text(json.dumps(config | {'max_files': 1}))
        settings = replace(settings, materials_path=path)
    else:
        settings = replace(settings, max_images=1)
    image = io.BytesIO()
    Image.new('RGB', (20, 20), 'white').save(image, format='PNG')
    client = Mock(spec=FeishuClient, settings=settings)
    client.download_attachment.return_value = image.getvalue()
    bundle = collect_materials({'fldA': [{'file_token': 'a', 'name': '素材.png'}],
                                'fldB': [{'file_token': 'b', 'name': '说明.png'}]}, None, 'recA', client)
    assert len(bundle.images) == 1 and len(bundle.skipped_attachments) == 1
    assert not bundle.issues
    assert client.download_attachment.call_count == (1 if limit == 'files' else 2)


@pytest.mark.parametrize("link", ["https://evil.example/sheets/abc", "https://kocotree.feishu.cn.evil.example/sheets/abc",
                                  "http://kocotree.feishu.cn/sheets/abc", "https://user@kocotree.feishu.cn/sheets/abc",
                                  "https://kocotree.feishu.cn/docx/abc", "https://127.0.0.1/sheets/abc"])
def test_sheet_link_rejects_unapproved_urls(config, link):
    with pytest.raises(ValueError):
        sheet_token(link, config)


def test_sheet_link_and_shared_image_limit(settings):
    feishu = Mock(spec=FeishuClient, settings=replace(settings, max_images=1))
    feishu.export_sheet.return_value = workbook_bytes()
    feishu.download_attachment.return_value = workbook_bytes()
    bundle = collect_materials({"fldTest": [{"file_token": "excel", "name": "说明.xlsx"}]},
                               {"link": "https://kocotree.feishu.cn/sheets/abc?sheet=123"}, "recTest", feishu)
    assert feishu.export_sheet.call_args.args[0] == "abc"
    assert len(bundle.images) == 1 and any("图片数量" in issue for issue in bundle.issues)


def test_export_poll_download_and_failure(settings, monkeypatch):
    feishu = FeishuClient(settings)
    monkeypatch.setattr(feishu, "_access_token", lambda: "secret")
    monkeypatch.setattr("app.feishu.time.sleep", Mock())
    request = Mock(side_effect=[{"data": {"ticket": "task"}}, {"data": {"result": {"job_status": 2}}},
                               {"data": {"result": {"job_status": 0, "file_token": "file", "file_size": 4}}}])
    monkeypatch.setattr(feishu, "_request", request)
    download = Mock(return_value=b"xlsx")
    monkeypatch.setattr(feishu, "_download", download)
    assert feishu.export_sheet("sheet", 100, 3, 1) == b"xlsx"
    assert request.call_args.kwargs["params"] == {"token": "sheet"}
    for status in (2, 110):
        request.side_effect = [{"data": {"ticket": "task"}}, {"data": {"result": {"job_status": status}}}]
        with pytest.raises(RuntimeError):
            feishu.export_sheet("sheet", 100, 1, 1)
    request.side_effect = [{"data": {"ticket": "task"}}, {"data": {"result": {"job_status": 0, "file_size": 101}}}]
    with pytest.raises(ValueError, match="大小"):
        feishu.export_sheet("sheet", 100, 1, 1)


def test_multiple_visible_sheets_survive_image_context(settings, config):
    """有图片的工作表不能影响后续工作表的解析配置。"""
    workbook = openpyxl.load_workbook(io.BytesIO(workbook_bytes()))
    extra = workbook.create_sheet("第二个工作表")
    extra["A1"] = "后续需求内容"
    output = io.BytesIO()
    workbook.save(output)
    bundle = MaterialBundle(settings, config)
    read_document(output.getvalue(), ".xlsx", "多页签.xlsx", bundle)
    assert not bundle.issues
    assert any("后续需求内容" in part["text"] for part in bundle.texts)
